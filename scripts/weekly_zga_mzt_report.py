from datetime import date, timedelta
from docx import Document
import requests
import re
import os

SEARCH_TERMS = [
    "zygotic genome activation",
    "ZGA",
    "maternal-to-zygotic transition",
    "maternal zygotic transition",
    "MZT",
]

PUBMED_QUERY = (
    '("zygotic genome activation" OR ZGA OR '
    '"maternal-to-zygotic transition" OR "maternal zygotic transition" OR MZT) '
    'AND (embryo OR "early embryo" OR embryogenesis OR oocyte)'
)


def clean_text(text):
    if not text:
        return ""
    return re.sub(r"\s+", " ", text).strip()


def get_pubmed_papers(start_date, end_date):
    search_url = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esearch.fcgi"
    fetch_url = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi"

    params = {
        "db": "pubmed",
        "term": PUBMED_QUERY,
        "retmode": "json",
        "retmax": 50,
        "mindate": start_date,
        "maxdate": end_date,
        "datetype": "pdat",
    }

    ids = requests.get(search_url, params=params, timeout=30).json()["esearchresult"]["idlist"]

    papers = []
    for pmid in ids:
        r = requests.get(
            fetch_url,
            params={"db": "pubmed", "id": pmid, "retmode": "xml"},
            timeout=30,
        )
        text = r.text

        title = re.search(r"<ArticleTitle>(.*?)</ArticleTitle>", text, re.S)
        abstract = re.findall(r"<AbstractText.*?>(.*?)</AbstractText>", text, re.S)
        journal = re.search(r"<Title>(.*?)</Title>", text, re.S)
        doi = re.search(r'<ArticleId IdType="doi">(.*?)</ArticleId>', text, re.S)

        papers.append(
            {
                "source": "PubMed",
                "title": clean_text(title.group(1)) if title else "No title found",
                "abstract": clean_text(" ".join(abstract)),
                "journal": clean_text(journal.group(1)) if journal else "PubMed",
                "date": str(end_date),
                "doi": clean_text(doi.group(1)) if doi else "",
                "link": f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/",
            }
        )

    return papers


def get_biorxiv_papers(start_date, end_date):
    url = f"https://api.biorxiv.org/details/biorxiv/{start_date}/{end_date}"
    data = requests.get(url, timeout=30).json()

    papers = []
    for item in data.get("collection", []):
        combined_text = f"{item.get('title', '')} {item.get('abstract', '')}".lower()
        if any(term.lower() in combined_text for term in SEARCH_TERMS):
            papers.append(
                {
                    "source": "bioRxiv",
                    "title": clean_text(item.get("title")),
                    "abstract": clean_text(item.get("abstract")),
                    "journal": "bioRxiv",
                    "date": item.get("date", ""),
                    "doi": item.get("doi", ""),
                    "link": f"https://doi.org/{item.get('doi')}" if item.get("doi") else "",
                }
            )

    return papers


def deduplicate(papers):
    seen = set()
    unique = []

    for paper in papers:
        key = paper.get("doi") or paper.get("title", "").lower()
        key = clean_text(key).lower()

        if key and key not in seen:
            seen.add(key)
            unique.append(paper)

    return unique


def make_simple_summary(abstract):
    if not abstract:
        return "No abstract was available. Review the linked paper manually."

    sentences = re.split(r"(?<=[.!?])\s+", abstract)
    return " ".join(sentences[:4])


def write_docx(papers, output_path, start_date, end_date):
    doc = Document()
    doc.add_heading("Weekly ZGA/MZT Paper Summary", level=1)

    doc.add_paragraph(f"Date range: {start_date} to {end_date}")
    doc.add_paragraph(f"Number of papers found: {len(papers)}")

    if not papers:
        doc.add_paragraph("No new matching papers were found this week.")

    for i, paper in enumerate(papers, start=1):
        doc.add_heading(f"{i}. {paper['title']}", level=2)
        doc.add_paragraph(f"Source: {paper['source']}")
        doc.add_paragraph(f"Journal/server: {paper['journal']}")
        doc.add_paragraph(f"Date: {paper['date']}")
        doc.add_paragraph(f"DOI: {paper['doi']}")
        doc.add_paragraph(f"Link: {paper['link']}")

        doc.add_heading("Abstract", level=3)
        doc.add_paragraph(paper["abstract"] or "No abstract available.")

        doc.add_heading("Brief summary", level=3)
        doc.add_paragraph(make_simple_summary(paper["abstract"]))

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)


def main():
    end_date = date.today()
    start_date = end_date - timedelta(days=7)

    papers = []
    papers.extend(get_pubmed_papers(start_date, end_date))
    papers.extend(get_biorxiv_papers(start_date, end_date))

    papers = deduplicate(papers)

    output_path = f"reports/ZGA_MZT_weekly_summary_{end_date}.docx"
    write_docx(papers, output_path, start_date, end_date)

    print(f"Saved report to {output_path}")


if __name__ == "__main__":
    main()
